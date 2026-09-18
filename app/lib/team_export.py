"""Builds the "download team formation" spreadsheet an admin gets from the
Teams page -- one row per student, grouped by team, so it can be pasted
straight into whatever else the course responsible tracks rosters in."""

from __future__ import annotations

import io
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def build_teams_workbook(
    assignment_title: str,
    teams_with_members: List[dict],
    unassigned: List[dict],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Teams"

    headers = ["Team", "Team code", "Student name", "Student email"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    row_count = 1
    for entry in teams_with_members:
        team = entry["team"]
        members = entry["members"]
        if not members:
            ws.append([team["name"], team["short_code"], "", ""])
            row_count += 1
            continue
        for member in members:
            ws.append([team["name"], team["short_code"], member["name"], member["email"]])
            row_count += 1

    if unassigned:
        ws.append(["Unassigned", "", "", ""])
        for cell in ws[row_count + 1]:
            cell.font = Font(italic=True)
        row_count += 1
        for student in unassigned:
            ws.append(["", "", student["name"], student["email"]])
            row_count += 1

    ws.freeze_panes = "A2"
    widths = [22, 12, 24, 30]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
